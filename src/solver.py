"""执行输入校验、模型求解、结果写出和审计记录。

运行方式：python main.py solve
"""

from __future__ import annotations

import hashlib
import json
import tempfile
from pathlib import Path
from dataclasses import asdict

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .arguments import ArgumentParser
from .dispatch import (
    ModeResult,
    _emergency_events,
    interval_label,
    load_inputs,
    mode_interval_frame,
    mode_summary,
    solve_question1,
)
from .causal_scenarios import CausalScenarioFactory
from .model_config import ModelConfig
from .optimization import adjustment_cost
from .paths import PROCESSED_DATA_DIR, PROJECT_ROOT, RAW_DATA_DIR, TEMPLATE_DIR, project_path
from .stochastic import execute_controls, solve_stochastic_plan


def weighted_quantile(values, probability, quantile):
    order = np.argsort(values)
    return float(
        np.asarray(values)[order][
            min(len(order) - 1, np.searchsorted(np.cumsum(probability[order]), quantile))
        ]
    )


def source_hashes(root):
    files = sorted(RAW_DATA_DIR.glob("*.xlsx"))
    files += sorted(TEMPLATE_DIR.glob("*.xlsx"))
    files += sorted(PROCESSED_DATA_DIR.glob("*.csv"))
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}


def january_initialization(data, config):
    """生成一月初始化状态：储能保持空闲，购电计划取前一日正净负荷。

    1 月 1 日无历史日，购电计划取零，正净负荷由紧急购电补足。一月只用于初始化和
    校准，其费用不计入二月至十二月结果；该规则不代表一月最优运行策略。
    """
    state = config.storage_initial_kwh
    records = []
    for day in range(31):
        q = np.maximum(data.net_kwh[day - 1], 0.0) if day else np.zeros(144)
        result = execute_controls(data.net_kwh[day], q, np.zeros(144), np.zeros(144), state, config)
        for t in range(144):
            records.append(
                dict(
                    date=data.dates[day],
                    interval_index=t + 1,
                    interval_start=interval_label(t)[0],
                    interval_end=interval_label(t)[1],
                    actual_net_load_kwh=data.net_kwh[day, t],
                    original_plan_kwh=q[t],
                    used_plan_kwh=result.used_plan_kwh[t],
                    emergency_kwh=result.emergency_kwh[t],
                    curtail_kwh=result.curtail_kwh[t],
                    charge_kwh=0.0,
                    discharge_kwh=0.0,
                    storage_start_kwh=state,
                    storage_kwh=result.storage_kwh[t],
                )
            )
        state = float(result.storage_kwh[-1])
    return state, pd.DataFrame(records)


def run_mode(
    data,
    factory,
    config,
    initial,
    *,
    name,
    rolling,
    pv_information,
    variable_price,
    unknown_price=False,
    days=334,
    risk_weight=0.20,
    gate=True,
):
    """求解指定模式，并限制决策只能读取历史日和当前发布时间前的数据。"""
    dates = data.dates[31 : 31 + days]
    names = (
        "original",
        "accepted",
        "used",
        "emergency",
        "charge",
        "discharge",
        "curtail",
        "storage",
        "center",
        "risk",
    )
    arrays = {key: np.zeros((days, 144)) for key in names}
    prices = (
        data.variable_price[31 : 31 + days].copy()
        if variable_price
        else np.broadcast_to(data.fixed_price, (days, 144)).copy()
    )
    starts, rows, gate_rows, audit_rows, bands = np.zeros(days), [], [], [], []
    state = float(initial)
    for out, day in enumerate(range(31, 31 + days)):
        starts[out] = state
        bundle = factory.get(
            day,
            pv_information=pv_information,
            variable_price=variable_price,
            unknown_price=unknown_price,
        )
        base = solve_stochastic_plan(
            bundle.net, bundle.prices, bundle.probability, state, config, risk_weight=risk_weight
        )
        original, accepted = base.q.copy(), base.q.copy()
        current = base
        arrays["center"][out], arrays["risk"][out] = bundle.center, bundle.upper
        actual = data.net_kwh[day]
        bands.append(
            dict(
                date=dates[out],
                lower_coverage=float(np.mean(actual >= bundle.lower)),
                interval_coverage=float(
                    np.mean((actual >= bundle.lower) & (actual <= bundle.upper))
                ),
                average_band_width_kwh=float(np.mean(bundle.upper - bundle.lower)),
                calibration_days=bundle.calibration_count,
            )
        )
        max_residual, accepted_count, projection = 0.0, 0, 0.0
        epochs = (0, 36, 72, 108) if rolling else (0,)
        for release, slot in enumerate(epochs):
            if slot:
                bundle = factory.get(
                    day,
                    release,
                    pv_information=True,
                    variable_price=variable_price,
                    unknown_price=unknown_price,
                )
                keep = solve_stochastic_plan(
                    bundle.net,
                    bundle.prices,
                    bundle.probability,
                    state,
                    config,
                    risk_weight=risk_weight,
                    original=original[slot:],
                    fixed_q=accepted[slot:],
                )
                candidate = solve_stochastic_plan(
                    bundle.net,
                    bundle.prices,
                    bundle.probability,
                    state,
                    config,
                    risk_weight=risk_weight,
                    original=original[slot:],
                )
                gain = keep.scenario_cost - candidate.scenario_cost
                # 按完整历史日重采样，避免人为拼接时段场景。
                # 经验重采样下分位数只用于门控诊断，不构成非平稳序列的理论保证。
                rng = np.random.default_rng(20250911 + day * 4 + release)
                sample = rng.choice(
                    len(gain), (300, max(8, len(bundle.source_indices))), p=bundle.probability
                )
                lower = float(np.quantile(gain[sample].mean(axis=1), 0.05))
                expected = float(bundle.probability @ gain)
                old_risk = weighted_quantile(keep.scenario_emergency, bundle.probability, 0.9)
                new_risk = weighted_quantile(candidate.scenario_emergency, bundle.probability, 0.9)
                threshold = config.gate_min_saving_ratio * max(
                    float(bundle.probability @ keep.scenario_cost), 1.0
                )
                economical = lower > threshold
                emergency = (
                    old_risk > config.gate_emergency_threshold_kwh
                    and new_risk < (1 - config.gate_required_risk_reduction_ratio) * old_risk
                )
                changed = bool(np.max(np.abs(candidate.q - accepted[slot:])) > 1e-6)
                take = bool(changed and ((economical or emergency) if gate else True))
                current = candidate if take else keep
                if take:
                    accepted[slot:] = candidate.q
                    accepted_count += 1
                gate_rows.append(
                    dict(
                        model=name,
                        date=dates[out],
                        release_time=f"{slot // 6:02d}:00",
                        accepted=take,
                        expected_saving_yuan=expected,
                        bootstrap_lower_yuan=lower,
                        threshold_yuan=threshold,
                        old_tail_emergency_kwh=old_risk,
                        new_tail_emergency_kwh=new_risk,
                        reason="economic"
                        if take and economical
                        else "emergency"
                        if take and emergency
                        else "forced"
                        if take
                        else "keep",
                    )
                )
                audit_rows.append(
                    dict(
                        model=name,
                        date=dates[out],
                        release_time=f"{slot // 6:02d}:00",
                        kind="keep",
                        objective_yuan=keep.objective,
                        residual=keep.residual,
                        milp=keep.used_milp,
                        mip_gap=keep.mip_gap,
                        latest_source_date=data.dates[bundle.source_indices.max()],
                        scenario_count=len(bundle.probability),
                    )
                )
            audit_rows.append(
                dict(
                    model=name,
                    date=dates[out],
                    release_time=f"{slot // 6:02d}:00",
                    kind="selected",
                    objective_yuan=current.objective,
                    residual=current.residual,
                    milp=current.used_milp,
                    mip_gap=current.mip_gap,
                    latest_source_date=data.dates[bundle.source_indices.max()],
                    scenario_count=len(bundle.probability),
                )
            )
            end = min(slot + 36, 144) if rolling else 144
            length = end - slot
            result = execute_controls(
                actual[slot:end],
                accepted[slot:end],
                current.charge[:length],
                current.discharge[:length],
                state,
                config,
            )
            projection += float(
                np.abs(result.charge_kwh - current.charge[:length]).sum()
                + np.abs(result.discharge_kwh - current.discharge[:length]).sum()
            )
            for key, attribute in (
                ("used", "used_plan_kwh"),
                ("emergency", "emergency_kwh"),
                ("charge", "charge_kwh"),
                ("discharge", "discharge_kwh"),
                ("curtail", "curtail_kwh"),
                ("storage", "storage_kwh"),
            ):
                arrays[key][out, slot:end] = getattr(result, attribute)
            state = float(result.storage_kwh[-1])
            max_residual = max(max_residual, result.max_balance_residual_kwh)
        arrays["original"][out], arrays["accepted"][out] = original, accepted
        purchase_cost = float(prices[out] @ original)
        adjust_cost = adjustment_cost(accepted, original, prices[out], config)
        emergency_cost = float(
            config.emergency_price_multiplier * (prices[out] @ arrays["emergency"][out])
        )
        rows.append(
            dict(
                date=dates[out],
                plan_energy_kwh=float(original.sum()),
                used_plan_energy_kwh=float(arrays["used"][out].sum()),
                emergency_energy_kwh=float(arrays["emergency"][out].sum()),
                charge_energy_kwh=float(arrays["charge"][out].sum()),
                discharge_energy_kwh=float(arrays["discharge"][out].sum()),
                curtail_energy_kwh=float(arrays["curtail"][out].sum()),
                storage_start_kwh=starts[out],
                storage_end_kwh=state,
                plan_cost_yuan=purchase_cost,
                adjustment_cost_yuan=adjust_cost,
                emergency_cost_yuan=emergency_cost,
                total_cost_yuan=purchase_cost + adjust_cost + emergency_cost,
                forecast_mae_kwh=float(np.abs(actual - arrays["center"][out]).mean()),
                max_balance_residual_kwh=max_residual,
                gate_count=accepted_count,
                control_projection_kwh=projection,
            )
        )
        if (out + 1) % 25 == 0 or out + 1 == days:
            print(
                f"{name}：已完成 {out + 1}/{days} 日，累计费用={sum(x['total_cost_yuan'] for x in rows):,.2f} 元",
                flush=True,
            )
    result = ModeResult(
        name,
        dates,
        arrays["original"],
        arrays["accepted"],
        arrays["used"],
        arrays["emergency"],
        arrays["charge"],
        arrays["discharge"],
        arrays["curtail"],
        arrays["storage"],
        starts,
        arrays["center"],
        arrays["risk"],
        prices,
        pd.DataFrame(rows),
        pd.DataFrame(gate_rows),
    )
    return result, pd.DataFrame(audit_rows), pd.DataFrame(bands)


def verify_mode(mode, data, config):
    """逐时段复核物理约束与费用。"""
    net = data.net_kwh[31 : 31 + len(mode.dates)]
    q, x, h = mode.accepted_plan_kwh, mode.used_plan_kwh, mode.emergency_kwh
    c, d, s, e = mode.charge_kwh, mode.discharge_kwh, mode.curtail_kwh, mode.storage_kwh
    previous = np.column_stack((mode.storage_start_kwh, e[:, :-1]))
    balance = float(np.abs(x + h + d - c - s - net).max())
    dynamics = float(np.abs(e - previous - config.eta_charge * c + d / config.eta_discharge).max())
    continuous = (
        float(np.max(np.abs(mode.storage_start_kwh[1:] - e[:-1, -1]))) if len(e) > 1 else 0.0
    )
    checks = dict(
        max_balance_residual_kwh=balance,
        max_soc_residual_kwh=dynamics,
        max_day_boundary_residual_kwh=continuous,
        max_simultaneous_kwh=float(np.minimum(c, d).max()),
        min_soc_kwh=float(e.min()),
        max_soc_kwh=float(e.max()),
        max_flow_kwh=float(max(c.max(), d.max())),
        max_used_over_plan_kwh=float(np.maximum(x - q, 0.0).max()),
        max_spill_over_pv_surplus_kwh=float(np.maximum(s - np.maximum(-net, 0.0), 0.0).max()),
        initial_soc_kwh=float(mode.storage_start_kwh[0]),
        final_soc_kwh=float(e[-1, -1]),
        all_finite=all(np.isfinite(a).all() for a in (q, x, h, c, d, s, e)),
    )
    assert (
        max(
            balance,
            dynamics,
            continuous,
            checks["max_simultaneous_kwh"],
            checks["max_used_over_plan_kwh"],
            checks["max_spill_over_pv_surplus_kwh"],
        )
        < 1e-5
    ), checks
    assert (
        checks["all_finite"]
        and e.min() >= config.storage_min_kwh - 1e-5
        and e.max() <= config.storage_max_kwh + 1e-5
    )
    assert checks["max_flow_kwh"] <= config.max_interval_energy_kwh + 1e-5
    assert min(a.min() for a in (q, x, h, c, d, s)) >= -1e-5
    recomputed = np.sum(mode.price_yuan_per_kwh * mode.original_plan_kwh, axis=1)
    recomputed += config.emergency_price_multiplier * np.sum(mode.price_yuan_per_kwh * h, axis=1)
    recomputed += [
        adjustment_cost(a, o, p, config)
        for a, o, p in zip(q, mode.original_plan_kwh, mode.price_yuan_per_kwh, strict=False)
    ]
    checks["max_cost_residual_yuan"] = float(
        np.abs(recomputed - mode.daily.total_cost_yuan.to_numpy()).max()
    )
    assert checks["max_cost_residual_yuan"] < 1e-5
    return checks


def write_workbook(root, output, template, mode=None, q1=None):
    """复制结果模板并写入标准时间标签与求解结果。

    输出全部日期、每日六个储能区块和所有连续紧急购电事件。原始模板保持只读，
    合计公式使用工作表单元格引用。
    """
    workbook = openpyxl.load_workbook(TEMPLATE_DIR / template)
    mapping = []
    plan_sheets = (
        workbook.worksheets[:2]
        if mode is not None and "调整" in workbook.worksheets[1].title
        else workbook.worksheets[:1]
    )
    for index, sheet in enumerate(plan_sheets):
        if q1 is not None:
            for t in range(144):
                mapping.append(
                    dict(
                        template=template,
                        sheet=sheet.title,
                        cell=f"A{t + 2}",
                        original_label=str(sheet.cell(t + 2, 1).value),
                        canonical_label="—".join(interval_label(t)),
                    )
                )
                sheet.cell(t + 2, 1, "—".join(interval_label(t)))
                sheet.cell(t + 2, 2, float(q1.purchase_kwh[t]))
            continue
        for t in range(144):
            mapping.append(
                dict(
                    template=template,
                    sheet=sheet.title,
                    cell=sheet.cell(1, t + 2).coordinate,
                    original_label=str(sheet.cell(1, t + 2).value),
                    canonical_label="—".join(interval_label(t)),
                )
            )
            sheet.cell(1, t + 2, "—".join(interval_label(t)))
        values = mode.original_plan_kwh if index == 0 else mode.accepted_plan_kwh
        for i, date in enumerate(mode.dates):
            sheet.cell(i + 2, 1, date.to_pydatetime()).number_format = "yyyy-mm-dd"
            for t in range(144):
                sheet.cell(i + 2, t + 2, float(values[i, t]))
            sheet.cell(i + 2, 146, f"=SUM(B{i + 2}:EO{i + 2})")
            cost = mode.daily.iloc[i].plan_cost_yuan + (
                mode.daily.iloc[i].adjustment_cost_yuan if index else 0.0
            )
            sheet.cell(i + 2, 147, float(cost))
        if sheet.max_row > len(mode.dates) + 1:
            sheet.delete_rows(len(mode.dates) + 2, sheet.max_row - len(mode.dates) - 1)
    if q1 is not None:
        sheet = workbook.worksheets[1]
        for block in range(6):
            sheet.cell(block + 2, 2, float(q1.charge_kwh[block * 24 : (block + 1) * 24].sum()))
            sheet.cell(block + 2, 3, float(q1.discharge_kwh[block * 24 : (block + 1) * 24].sum()))
        sheet.cell(2, 5, 6000.0)
        sheet.cell(3, 5, float(q1.storage_kwh[-1]))
    else:
        sheet = workbook.worksheets[len(plan_sheets)]
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(
            [
                "日期",
                "时段（起点—终点）",
                "充电量/kWh",
                "放电量/kWh",
                "起始储电量/kWh",
                "结束储电量/kWh",
            ]
        )
        for i, date in enumerate(mode.dates):
            for block in range(6):
                start, end = block * 24, (block + 1) * 24
                sheet.append(
                    [
                        date.to_pydatetime(),
                        f"{block * 4:02d}:00—{(block + 1) * 4:02d}:00",
                        float(mode.charge_kwh[i, start:end].sum()),
                        float(mode.discharge_kwh[i, start:end].sum()),
                        float(
                            mode.storage_start_kwh[i]
                            if not start
                            else mode.storage_kwh[i, start - 1]
                        ),
                        float(mode.storage_kwh[i, end - 1]),
                    ]
                )
        sheet = workbook.worksheets[len(plan_sheets) + 1]
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(["日期", "紧急购电时段（连续段）", "紧急购电量/kWh"])
        for i, date in enumerate(mode.dates):
            events = _emergency_events(mode.emergency_kwh[i])
            for label, energy in events or [("无", 0.0)]:
                sheet.append([date.to_pydatetime(), label, float(energy)])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "B2" if sheet in plan_sheets and mode else "A2"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 36
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(
                    name="Microsoft YaHei",
                    size=10,
                    color="008000" if isinstance(cell.value, float | int) else "000000",
                )
                if isinstance(cell.value, float | int):
                    cell.number_format = '#,##0.000;[Red](#,##0.000);"—"'
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.number_format = '#,##0.000;[Red](#,##0.000);"—"'
            if mode:
                row[0].number_format = "yyyy-mm-dd"
        sheet.column_dimensions["A"].width = 23 if q1 else 15
        if sheet in plan_sheets and mode:
            for col in range(2, 148):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 17
        else:
            for col in range(2, sheet.max_column + 1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 27
        sheet.sheet_view.zoomScale = 85
    info = workbook.create_sheet("口径说明")
    notes = [
        ("说明", "本文件为修正时间标签的计算结果副本；未改变原模板文件。"),
        (
            "时间",
            "144时段按00:00—24:00；原模板部分标签偏移10分钟，完整映射见template_time_mapping.csv。",
        ),
        ("购电", "计划电量均为10分钟电量kWh，不是kW；未使用计划电仍计费。"),
        ("储能", "每侧效率0.9；SOC范围1200—10800kWh；母线侧充/放电功率各不超过5000kW。"),
        (
            "调整",
            "按最终计划相对原始计划的净调增、净调减计费；非逐次交易累加；退款系数见运行参数。",
        ),
        (
            "初始化",
            "1月采用留存6000kWh的闲置储能校准策略，2月1日继承1月31日实态；1月费用不计考核。",
        ),
        ("数据范围", "仅输出已运行日期；完整年度默认2025-02-01至2025-12-31。"),
        (
            "结果性质",
            "Q1确定性LP；其余为有限场景Wasserstein+CVaR与共同储能动作策略，不是无限支持全局最优。",
        ),
        (
            "格式兼容",
            "充放电/紧急购电示例行已展开为全部日期/全部连续段；是否允许修正模板需赛方确认。",
        ),
    ]
    for row in notes:
        info.append(row)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 110
    for row in info:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(name="Microsoft YaHei", size=10)
        info.row_dimensions[row[0].row].height = 34
    path = output / template
    workbook.save(path)
    return path, mapping


def verify_workbook(path, mode=None, q1=None):
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook.worksheets[0]
    if q1 is not None:
        assert sheet.cell(2, 1).value == "00:00—00:10" and sheet.cell(145, 1).value == "23:50—24:00"
        np.testing.assert_allclose(
            [sheet.cell(t + 2, 2).value for t in range(144)], q1.purchase_kwh, atol=1e-8
        )
    else:
        assert sheet.max_row == len(mode.dates) + 1 and sheet.cell(1, 2).value == "00:00—00:10"
        plan_count = 2 if "调整" in workbook.worksheets[1].title else 1
        for idx in range(plan_count):
            ps = workbook.worksheets[idx]
            expected = mode.original_plan_kwh if idx == 0 else mode.accepted_plan_kwh
            np.testing.assert_allclose(
                [[ps.cell(i + 2, t + 2).value for t in range(144)] for i in range(len(mode.dates))],
                expected,
                atol=1e-8,
            )
            for i in range(len(mode.dates)):
                assert ps.cell(i + 2, 146).value == f"=SUM(B{i + 2}:EO{i + 2})"
        bs = workbook.worksheets[plan_count]
        assert bs.max_row == 1 + 6 * len(mode.dates)
        np.testing.assert_allclose(
            sum(float(bs.cell(r, 3).value) for r in range(2, bs.max_row + 1)),
            mode.charge_kwh.sum(),
            atol=1e-6,
        )
        es = workbook.worksheets[plan_count + 1]
        np.testing.assert_allclose(
            sum(float(es.cell(r, 3).value) for r in range(2, es.max_row + 1)),
            mode.emergency_kwh.sum(),
            atol=1e-5,
        )
    for ws in workbook.worksheets:
        assert not any(c.data_type == "e" for row in ws for c in row)
    return dict(file=path.name, verified=True, sheets=workbook.sheetnames)


def _run(output):
    parser = ArgumentParser(description="执行购电计划优化并生成交付结果")
    parser.add_argument("--output", help="结果目录；默认按全部问题或单问分别保存")
    parser.add_argument("--question", choices=["all", "1", "2", "3", "4"], default="all", help="求解范围")
    parser.add_argument(
        "--days", type=int, default=334, help="求解天数；小于 334 时结果标记为部分运行"
    )
    parser.add_argument("--scenarios", type=int, default=8, help="代表场景数量")
    parser.add_argument("--risk-weight", type=float, default=0.20, help="尾部风险权重")
    parser.add_argument("--radius", type=float, default=0.15, help="分布鲁棒半径")
    parser.add_argument("--refund", type=float, default=0.0, help="调减购电的退款比例")
    parser.add_argument(
        "--price-information", choices=["known", "unknown"], default="known",
        help="电价信息：known 为日前已知，unknown 为预测电价",
    )
    parser.add_argument(
        "--ablations", action="store_true", help="增加问题 3 和问题 4-3 的日内不更新对照组"
    )
    args = parser.parse_args()
    if args.output is None:
        args.output = "result"
    if not 1 <= args.days <= 334:
        parser.error("days 必须在 1 到 334 之间")
    root = PROJECT_ROOT
    destination = project_path(args.output).resolve()
    result_root = root / "result"
    if not destination.is_relative_to(result_root):
        parser.error("交付目录必须位于 result 下")
    if args.days < 334 and args.question != "1" and destination == result_root:
        parser.error("部分日期运行请指定 --output result/check，避免覆盖正式结果")
    config = ModelConfig(
        robustness_radius=args.radius,
        down_refund_ratio=args.refund,
        main_price_information_case="known_day_ahead"
        if args.price_information == "known"
        else "historically_forecast",
    )
    config.validate()
    hashes = source_hashes(root)
    data = load_inputs(root, config)
    factory = CausalScenarioFactory(data, config, args.scenarios)
    initial, january = january_initialization(data, config)
    january.to_csv(output / "january_initialization.csv", index=False, encoding="utf-8-sig")
    q1, q1frame = solve_question1(data, config)
    if args.question in {"all", "1"}:
        q1frame.to_csv(output / "question1_intervals.csv", index=False, encoding="utf-8-sig")
    specifications = [
        ("Q2", False, False, False, "result2.xlsx"),
        ("Q3", True, True, False, "result3.xlsx"),
        ("Q4-2", False, False, True, "result4-2.xlsx"),
        ("Q4-3", True, True, True, "result4-3.xlsx"),
    ]
    if args.ablations:
        specifications += [
            ("Q3-no-update", False, True, False, None),
            ("Q4-3-no-update", False, True, True, None),
        ]
    if args.question != "all":
        prefix = {"1": "Q1", "2": "Q2", "3": "Q3", "4": "Q4-"}[args.question]
        specifications = [item for item in specifications if item[0].startswith(prefix)]
    modes = {}
    checks = {}
    all_audits = []
    all_bands = []
    mappings = []
    workbooks = []
    if args.question in {"all", "1"}:
        path, mapping = write_workbook(root, output, "result1.xlsx", q1=q1)
        mappings += mapping
        workbooks.append(verify_workbook(path, q1=q1))
    for name, rolling, pv_information, variable_price, template in specifications:
        mode, audits, bands = run_mode(
            data,
            factory,
            config,
            initial,
            name=name,
            rolling=rolling,
            pv_information=pv_information,
            variable_price=variable_price,
            unknown_price=args.price_information == "unknown",
            days=args.days,
            risk_weight=args.risk_weight,
        )
        checks[name] = verify_mode(mode, data, config)
        modes[name] = mode
        bands.insert(0, "model", name)
        all_audits.append(audits)
        all_bands.append(bands)
        mode.daily.to_csv(output / f"{name}_daily.csv", index=False, encoding="utf-8-sig")
        mode_interval_frame(mode, data.net_kwh[31 : 31 + args.days]).to_csv(
            output / f"{name}_intervals.csv", index=False, encoding="utf-8-sig"
        )
        mode.gates.to_csv(output / f"{name}_gates.csv", index=False, encoding="utf-8-sig")
        if template:
            path, mapping = write_workbook(root, output, template, mode=mode)
            mappings += mapping
            workbooks.append(verify_workbook(path, mode=mode))
    summary = pd.DataFrame([mode_summary(mode) for mode in modes.values()])
    summary.to_csv(output / "model_summary.csv", index=False, encoding="utf-8-sig")
    (pd.concat(all_audits, ignore_index=True) if all_audits else pd.DataFrame()).to_csv(
        output / "solver_audit.csv", index=False, encoding="utf-8-sig"
    )
    (pd.concat(all_bands, ignore_index=True) if all_bands else pd.DataFrame()).to_csv(
        output / "forecast_coverage.csv", index=False, encoding="utf-8-sig"
    )
    pd.DataFrame(mappings).to_csv(
        output / "template_time_mapping.csv", index=False, encoding="utf-8-sig"
    )
    assert hashes == source_hashes(root), "运行期间原始数据、模板或处理后数据发生变化"
    baseline = float(
        data.fixed_price
        @ np.maximum(
            data.typical.load_energy_kwh.to_numpy()
            - data.typical.pv_forecast_energy_kwh.to_numpy(),
            0,
        )
    )
    info = dict(
        version="finite_support_dro_v2",
        partial=args.days < 334,
        parameters=asdict(config),
        solver_parameters=vars(args),
        question1=dict(
            baseline_cost_yuan=baseline,
            cost_yuan=float(data.fixed_price @ q1.purchase_kwh),
            max_balance_residual_kwh=q1.max_balance_residual_kwh,
            simultaneous_kwh=q1.simultaneous_flow_kwh,
        ),
        models=summary.to_dict(orient="records"),
        physical_checks=checks,
        workbook_checks=workbooks,
        source_hashes=hashes,
        limitations=[
            "模型采用有限缩减支撑集，不代表无限支撑 Wasserstein 模型",
            "各场景共用储能控制，不是完整多阶段补救模型",
            "结果属于历史回放，不是独立留出检验，也不保证区间覆盖率",
            "调减退款系数和电价已知假设需要由题目负责人确认",
            "工作簿公式缓存需要使用表格软件重新计算",
        ],
    )
    (output / "solution_summary.json").write_text(
        json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    from .reporting import create_figures, create_question_figures
    from .delivery import publish_results
    from .paths import figure_directory

    if args.question == "all":
        captions = create_figures(output, data, q1frame, info, modes, config)
    else:
        captions = create_question_figures(output, data, q1frame, info, modes, config)
    publish_results(output, destination, info, figure_directory(output),
                    update_assets=args.question == "all" and args.days == 334 and destination == result_root)
    print(f"最终结果已保存：{destination}", flush=True)


def main():
    """中间计算放入系统临时目录，只发布最终交付文件。"""
    with tempfile.TemporaryDirectory(prefix="sxjm-solve-") as directory:
        _run(Path(directory))


if __name__ == "__main__":
    main()
